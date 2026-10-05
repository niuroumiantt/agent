import json
import sqlite3

import pytest
from docx import Document
from openpyxl import load_workbook

from glocal_agent.config import Settings, load_settings
from glocal_agent.reports import write_reports
from glocal_agent.store import Store


def test_no_other_application_credential_inheritance(tmp_path, monkeypatch):
    monkeypatch.setenv("AGENT_CONFIG", str(tmp_path / "absent.json"))
    for key in ("AGENT_BASE_URL", "AGENT_API_KEY", "AGENT_PROVIDER", "AGENT_MODEL"):
        monkeypatch.delenv(key, raising=False)
    monkeypatch.setenv("DGX_GATEWAY_URL", "https://other-app.example")
    monkeypatch.setenv("DGX_API_KEY", "other-app-private-key")
    monkeypatch.setenv("OLLAMA_BASE_URL", "http://other-device:11434")
    data = tmp_path / "private"
    root = tmp_path / "Downloads"
    root.mkdir()
    settings = load_settings(str(root), str(data))
    assert settings.base_url == ""
    assert settings.api_key == ""
    assert settings.public()["configured"] is False


def test_data_directory_cannot_pollute_inputs(tmp_path):
    with pytest.raises(ValueError, match="之外"):
        Settings(tmp_path, tmp_path / "reports").validate()


def test_store_single_instance_and_restart_marks_interrupted(tmp_path):
    first = Store(tmp_path / "data")
    job_id = first.create("test", "local-user", [])
    first.update(job_id, "running")
    with pytest.raises(ValueError, match="已有"):
        Store(tmp_path / "data")
    assert first.get(job_id)["status"] == "running"
    first.close()
    resumed = Store(tmp_path / "data")
    assert resumed.get(job_id)["status"] == "interrupted"
    assert resumed.get(job_id)["error"]
    assert resumed.path.stat().st_mode & 0o777 == 0o600
    resumed.close()


def test_existing_task_database_migrates_and_retains_results(tmp_path):
    directory = tmp_path / "data"
    directory.mkdir()
    with sqlite3.connect(directory / "tasks.sqlite3") as db:
        db.execute("""CREATE TABLE jobs (
            id TEXT PRIMARY KEY, status TEXT, created_at TEXT, updated_at TEXT,
            instruction TEXT, actor TEXT, payload TEXT, result TEXT, artifacts TEXT, error TEXT
        )""")
        db.execute("INSERT INTO jobs VALUES (?,?,?,?,?,?,?,?,?,?)", (
            "old", "completed", "2026", "2026", "read", "user", "[]",
            '{"summary":"old result"}', "[]", None,
        ))
    store = Store(directory)
    try:
        assert store.get("old")["result"]["summary"] == "old result"
        assert store.get("old")["progress"] == {}
        new_id = store.create("read", "user", [])
        store.update(new_id, "running", progress={"files_done": 1})
        store.update(new_id, "partial", result={"summary": "partial"})
        assert store.get(new_id)["progress"]["files_done"] == 1
    finally:
        store.close()


def test_office_exports_handle_controls_and_formula_like_untrusted_text(tmp_path):
    original_quote = '=HYPERLINK("https://example.invalid")\f Quantity 32'
    result = {
        "summary": "摘要\x00内容", "documents": [],
        "facts": [{"source_id": "S1", "locator": "page 1", "quote": original_quote,
                   "claim": "+1+1", "verified": True}],
        "recommendations": ["人工核对"], "warnings": [],
        "sources": [{"source_id": "S1", "file": {"relative_path": "=untrusted.xlsx"},
                     "status": "ok", "sha256": "a" * 64, "warnings": []}],
    }
    artifacts = write_reports(tmp_path, "test-job", result)
    assert len(artifacts) == 4
    folder = tmp_path / "artifacts/test-job"
    original = json.loads((folder / "report.json").read_text())
    assert original["facts"][0]["quote"] == original_quote
    assert Document(folder / "report.docx").paragraphs
    book = load_workbook(folder / "facts.xlsx")
    assert book["候选事实"]["D2"].data_type == "s"
    assert book["候选事实"]["E2"].data_type == "s"
    assert book["来源与限制"]["B2"].data_type == "s"
    assert (folder / "report.docx").stat().st_mode & 0o777 == 0o600
