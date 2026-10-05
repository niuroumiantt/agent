from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from openpyxl import Workbook

NOTICE = "AI 分析建议。原文命中仅证明引用存在，不证明模型解释、金额或业务结论正确。"


def office_text(text: str) -> str:
    """Keep XML-valid display text; the JSON report retains exact original quotes."""
    return "".join(
        character if (
            character in "\t\r\n" or 0x20 <= ord(character) <= 0xD7FF
            or 0xE000 <= ord(character) <= 0xFFFD or 0x10000 <= ord(character) <= 0x10FFFF
        ) else "�"
        for character in text
    )


def write_reports(directory: Path, job_id: str, result: dict) -> list[dict]:
    output = directory / "artifacts" / job_id
    output.mkdir(parents=True, exist_ok=True, mode=0o700)
    output.chmod(0o700)
    lines = ["# 文件分析报告", "", NOTICE, "", result["summary"], "", "## 资料"]
    doc = Document()
    doc.add_heading("文件分析报告", 0)
    doc.add_paragraph(NOTICE)
    doc.add_paragraph(office_text(result["summary"]))
    source_names = {
        source["source_id"]: source["file"]["relative_path"] for source in result["sources"]
    }
    for source in result["sources"]:
        line = f'{source["source_id"]} · {source["file"]["relative_path"]} · {source["status"]}'
        lines.extend(["", line, f'SHA-256: {source["sha256"]}'])
        doc.add_paragraph(office_text(line))
        doc.add_paragraph(f'SHA-256: {source["sha256"]}')
        for warning in source["warnings"]:
            lines.append("注意：" + warning)
            doc.add_paragraph(office_text("注意：" + warning))
    doc.add_heading("文件归纳", 1)
    lines.extend(["", "## 文件归纳"])
    for item in result["documents"]:
        line = f'{item["source_id"]} · {item["category"]}: {item["summary"]}'
        lines.extend(["", line])
        doc.add_paragraph(office_text(line))
    doc.add_heading("候选事实与原文", 1)
    lines.extend(["", "## 候选事实与原文"])
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "候选事实"
    sheet.append(["来源", "文件", "位置", "原文引用", "AI解释", "原文命中"])
    for item in result["facts"]:
        flag = "原文命中" if item["verified"] else "未命中，请核对"
        line = f'{item["source_id"]} / {item["locator"]} / {flag}: {item["claim"]}'
        lines.extend(["", line, "原文：" + item["quote"]])
        doc.add_paragraph(office_text(line))
        doc.add_paragraph(office_text("原文：" + item["quote"]))
        values = [
            item["source_id"], source_names.get(item["source_id"], "未知来源"),
            item["locator"], item["quote"], item["claim"], flag,
        ]
        row = sheet.max_row + 1
        for column, value in enumerate(values, 1):
            cell = sheet.cell(row, column, office_text(value))
            cell.data_type = "s"  # 文件和模型文字不能成为可执行的 Excel 公式。
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column, width in {"A": 10, "B": 35, "C": 25, "D": 70, "E": 60, "F": 24}.items():
        sheet.column_dimensions[column].width = width
    sources_sheet = workbook.create_sheet("来源与限制")
    sources_sheet.append(["来源", "文件", "SHA-256", "读取状态", "限制"])
    for source in result["sources"]:
        row = sources_sheet.max_row + 1
        values = [
            source["source_id"], source["file"]["relative_path"], source["sha256"],
            source["status"], "；".join(source["warnings"]),
        ]
        for column, value in enumerate(values, 1):
            sources_sheet.cell(row, column, office_text(value)).data_type = "s"
    doc.add_heading("行动建议", 1)
    lines.extend(["", "## 行动建议"])
    for item in result["recommendations"]:
        lines.extend(["", "- " + item])
        doc.add_paragraph(office_text(item))
    for warning in result["warnings"]:
        lines.extend(["", "注意：" + warning])
        doc.add_paragraph(office_text("注意：" + warning))
    (output / "report.md").write_text(office_text("\n".join(lines)), encoding="utf-8")
    (output / "report.json").write_text(
        json.dumps(result, ensure_ascii=True, indent=2), encoding="utf-8"
    )
    doc.save(output / "report.docx")
    workbook.save(output / "facts.xlsx")
    names = ["report.md", "report.json", "report.docx", "facts.xlsx"]
    for name in names:
        (output / name).chmod(0o600)
    return [{"name": name, "url": f"/api/jobs/{job_id}/artifacts/{name}"} for name in names]
