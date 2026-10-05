import hashlib
import json
import threading
import time

import httpx
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook

from glocal_agent.app import create_app
from glocal_agent.config import Settings
from glocal_agent.model import ModelError, analyze


def setup(tmp_path, analyzer):
    root = tmp_path / "Downloads"
    root.mkdir()
    (root / "po.txt").write_text("Quantity: 32", encoding="utf-8")
    app = create_app(
        Settings(root, tmp_path / "private", base_url="https://unused.example"), analyzer
    )
    return app, root, {"X-Agent-Session": app.state.session_token}


def finished(client, job_id, headers):
    deadline = time.monotonic() + 3
    while time.monotonic() < deadline:
        task = client.get(f"/api/jobs/{job_id}", headers=headers).json()
        if task["status"] not in {"queued", "running"}:
            return task
        time.sleep(0.01)
    raise AssertionError("Synthetic job did not finish")


def fake_analysis(settings, instruction, sources):
    return {
        "summary": "需要核对的采购材料", "documents": [],
        "facts": [{"source_id": "S1", "locator": "line 1", "quote": "Quantity: 32",
                   "claim": "32", "verified": True}],
        "recommendations": ["核对后另存"], "warnings": [], "model": settings.model,
    }


def test_selected_file_to_persisted_reports_without_changing_original(tmp_path):
    app, root, headers = setup(tmp_path, fake_analysis)
    original = (root / "po.txt").read_bytes()
    with TestClient(app) as client:
        assert "__SESSION_TOKEN__" not in client.get("/").text
        scan = client.post("/api/scan", json={}, headers=headers).json()
        file_id = scan["files"][0]["id"]
        preview = client.get(f"/api/files/{file_id}/preview", headers=headers).json()
        assert preview["sha256"] == hashlib.sha256(original).hexdigest()
        response = client.post("/api/jobs", json={
            "file_ids": [file_id], "instruction": "识别采购资料"
        }, headers=headers)
        assert response.status_code == 202
        task = finished(client, response.json()["id"], headers)
        assert task["status"] == "completed"
        assert task["result"]["sources"][0]["sha256"] == preview["sha256"]
        assert len(task["artifacts"]) == 4
        artifact = client.get(task["artifacts"][2]["url"], headers=headers)
        assert artifact.status_code == 200
        assert artifact.content.startswith(b"PK")
        assert client.get(task["artifacts"][2]["url"]).status_code == 403
        assert client.get(f'/api/jobs/{task["id"]}/artifacts/not-a-report',
                          headers=headers).status_code == 404
        history = client.get("/api/jobs", headers=headers).json()
        assert history["jobs"][0]["status"] == "completed"
    assert (root / "po.txt").read_bytes() == original
    assert list(root.iterdir()) == [root / "po.txt"]


def test_loopback_session_and_origin_boundaries(tmp_path):
    app, _, headers = setup(tmp_path, fake_analysis)
    with TestClient(app) as client:
        assert client.get("/api/config").status_code == 403
        assert client.get("/api/config", headers={"X-Agent-Session": "wrong"}).status_code == 403
        for origin in ("https://evil.example", "http://localhost:9999"):
            assert client.post("/api/scan", json={}, headers={
                **headers, "Origin": origin
            }).status_code == 403
        assert client.get("/api/config", headers={
            **headers, "Host": "evil.example"
        }).status_code == 400
        assert client.post("/api/scan", content="{}", headers=headers).status_code == 415
        assert client.get("/api/config", headers=headers).status_code == 200
        assert "api_key" not in client.get("/api/config", headers=headers).json()


def test_model_failure_is_visible_and_gate_is_released(tmp_path):
    def fail(settings, instruction, sources):
        raise ModelError("模型拒绝访问")

    app, _, headers = setup(tmp_path, fail)
    with TestClient(app) as client:
        scan = client.post("/api/scan", json={}, headers=headers).json()
        payload = {"file_ids": [scan["files"][0]["id"]], "instruction": "阅读"}
        for _ in range(2):
            response = client.post("/api/jobs", json=payload, headers=headers)
            assert response.status_code == 202
            task = finished(client, response.json()["id"], headers)
            assert task["status"] == "failed"
            assert task["error"] == "模型拒绝访问"
            assert task["result"] is None


def test_unsupported_material_never_sent_to_model(tmp_path):
    calls = []

    def should_not_call(*args):
        calls.append(args)
        raise AssertionError

    app, root, headers = setup(tmp_path, should_not_call)
    (root / "scan.png").write_bytes(b"image")
    with TestClient(app) as client:
        files = client.post("/api/scan", json={}, headers=headers).json()["files"]
        image = next(item for item in files if item["name"] == "scan.png")
        response = client.post("/api/jobs", json={
            "file_ids": [image["id"]], "instruction": "读图"
        }, headers=headers)
        assert response.status_code == 422
    assert not calls


def test_all_extracted_text_reaches_model_in_bounded_segments(tmp_path):
    captured = []

    def capture(settings, instruction, sources):
        if sources[0]["source_id"] != "S0":
            captured.extend(sources)
        return fake_analysis(settings, instruction, sources)

    app, root, headers = setup(tmp_path, capture)
    text = "a" * 18000 + "末尾条款：付款 90 天"
    (root / "po.txt").write_text(text)
    with TestClient(app) as client:
        file_id = client.post("/api/scan", json={}, headers=headers).json()["files"][0]["id"]
        response = client.post("/api/jobs", json={
            "file_ids": [file_id], "instruction": "阅读"
        }, headers=headers)
        task = finished(client, response.json()["id"], headers)
        assert task["status"] == "completed"
        assert task["result"]["sources"][0]["segments_done"] == 3
    assert "".join(block["text"] for source in captured for block in source["blocks"]) == text
    assert all(sum(len(block["text"]) for block in source["blocks"]) <= 8000
               for source in captured)
    assert all(source["status"] == "ok" for source in captured)


def test_more_than_six_files_and_recursive_selection(tmp_path):
    captured = []

    def capture(settings, instruction, sources):
        captured.append(sources[0]["source_id"])
        return {**fake_analysis(settings, instruction, sources), "facts": []}

    app, root, headers = setup(tmp_path, capture)
    folder = root / "contracts"
    folder.mkdir()
    for number in range(7):
        (folder / f"contract-{number}.txt").write_text(f"Price: {number}")
    with TestClient(app) as client:
        files = client.post("/api/scan", json={"recursive": True}, headers=headers).json()["files"]
        assert len(files) == 8
        response = client.post("/api/jobs", json={
            "file_ids": [file["id"] for file in files], "instruction": "比较所有合同"
        }, headers=headers)
        assert response.status_code == 202
        task = finished(client, response.json()["id"], headers)
        assert task["status"] == "completed"
        assert task["progress"]["files_done"] == 8
        assert len(task["result"]["sources"]) == 8
    assert set(captured) == {"S0", *(f"S{number}" for number in range(1, 9))}


def test_credentials_are_blocked_in_preview_and_never_reach_analyzer(tmp_path):
    captured = []

    def capture(settings, instruction, sources):
        captured.extend(sources)
        return fake_analysis(settings, instruction, sources)

    app, root, headers = setup(tmp_path, capture)
    (root / "company_credentials.csv").write_text("secret material")
    (root / "private.pem").write_text("private material")
    (root / "innocent.csv").write_text("user,aws_secret_access_key\nalice,synthetic-value")
    with TestClient(app) as client:
        files = client.post("/api/scan", json={}, headers=headers).json()["files"]
        for name in ("company_credentials.csv", "private.pem"):
            file = next(file for file in files if file["name"] == name)
            assert not file["supported"]
            response = client.get(f'/api/files/{file["id"]}/preview', headers=headers)
            assert response.status_code == 400
        innocent = next(file for file in files if file["name"] == "innocent.csv")
        preview = client.get(f'/api/files/{innocent["id"]}/preview', headers=headers).json()
        assert preview["extraction"]["status"] == "blocked"
        assert preview["extraction"]["blocks"] == []
        response = client.post("/api/jobs", json={
            "file_ids": [file["id"] for file in files], "instruction": "阅读"
        }, headers=headers)
        task = finished(client, response.json()["id"], headers)
        assert task["status"] == "partial"
        assert task["progress"]["files_skipped"] == 3
        assert task["progress"]["files_analyzed"] == 1
        assert "synthetic-value" not in str(task)
        assert "secret material" not in str(task)
        assert len(task["artifacts"]) == 4
    assert [source["file"]["name"] for source in captured] == ["po.txt"]


def test_cancellation_keeps_completed_segments_and_scan_snapshot(tmp_path):
    entered, release = threading.Event(), threading.Event()
    captured = []

    def wait(settings, instruction, sources):
        captured.append(sources)
        entered.set()
        assert release.wait(3)
        return fake_analysis(settings, instruction, sources)

    app, root, headers = setup(tmp_path, wait)
    (root / "po.txt").write_text("a" * 17000)
    with TestClient(app) as client:
        files = client.post("/api/scan", json={}, headers=headers).json()["files"]
        response = client.post("/api/jobs", json={
            "file_ids": [files[0]["id"]], "instruction": "阅读"
        }, headers=headers)
        job_id = response.json()["id"]
        assert entered.wait(3)
        try:
            assert client.post("/api/jobs", json={
                "file_ids": [files[0]["id"]], "instruction": "阅读"
            }, headers=headers).status_code == 409
            response = client.post(f"/api/jobs/{job_id}/cancel", json={}, headers=headers)
            assert response.status_code == 202
            client.post("/api/scan", json={}, headers=headers)
        finally:
            release.set()
        task = finished(client, job_id, headers)
        assert task["status"] == "cancelled"
        source = task["result"]["sources"][0]
        assert source["segments_done"] == 1
        assert source["segments_total"] == 3
        assert "segments_incomplete" in source["warnings"]
        assert len(task["artifacts"]) == 4
        assert client.get(task["artifacts"][0]["url"], headers=headers).status_code == 200
    assert len(captured) == 1


def test_changed_file_is_not_rebound_by_rescanning_during_job(tmp_path):
    entered, release = threading.Event(), threading.Event()
    captured = []

    def wait(settings, instruction, sources):
        captured.append(sources[0])
        entered.set()
        assert release.wait(3)
        return fake_analysis(settings, instruction, sources)

    app, root, headers = setup(tmp_path, wait)
    (root / "z-contract.txt").write_text("Original terms")
    with TestClient(app) as client:
        files = client.post("/api/scan", json={}, headers=headers).json()["files"]
        response = client.post("/api/jobs", json={
            "file_ids": [file["id"] for file in files], "instruction": "阅读"
        }, headers=headers)
        assert entered.wait(3)
        try:
            (root / "z-contract.txt").write_text("Replaced terms")
            client.post("/api/scan", json={}, headers=headers)
        finally:
            release.set()
        task = finished(client, response.json()["id"], headers)
        assert task["status"] == "partial"
        source = task["result"]["sources"][1]
        assert "rescan_required" in source["warnings"]
        assert not source["sha256"]
    assert len(captured) == 1


def test_office_batch_checks_original_cell_and_paragraph_quotes(tmp_path):
    seen = []

    def respond(request):
        payload = json.loads(request.content)
        source = json.loads(payload["messages"][1]["content"])["sources"][0]
        seen.append(source)
        block = source["blocks"][0]
        fact = {"source_id": source["source_id"], "locator": block["locator"],
                "quote": block["text"], "claim": "需要核对"}
        reply = {"summary": "合成资料归纳", "documents": [],
                 "facts": [] if source["source_id"] == "S0" else [fact],
                 "recommendations": []}
        return httpx.Response(200, json={"message": {"content": json.dumps(reply)}})

    with httpx.Client(transport=httpx.MockTransport(respond)) as model_client:
        def actual_analyzer(settings, instruction, sources):
            return analyze(settings, instruction, sources, client=model_client)

        app, root, headers = setup(tmp_path, actual_analyzer)
        doc = Document()
        doc.add_paragraph("Payment: 90 days")
        doc.save(root / "contract.docx")
        book = Workbook()
        book.active["A1"] = "PO quantity: 32"
        book.save(root / "po.xlsx")
        with TestClient(app) as client:
            files = client.post("/api/scan", json={}, headers=headers).json()["files"]
            response = client.post("/api/jobs", json={
                "file_ids": [file["id"] for file in files], "instruction": "核对合同与 PO"
            }, headers=headers)
            task = finished(client, response.json()["id"], headers)
            assert task["status"] == "completed"
            assert all(fact["verified"] for fact in task["result"]["facts"])
            assert {fact["locator"] for fact in task["result"]["facts"]} == {
                "paragraph 1", "'Sheet'!A1", "line 1",
            }
            assert len(task["result"]["facts"]) == 3
    assert seen[-1]["source_id"] == "S0"
