"""Bounded document extraction with source locations and explicit limitations.

This module reads bytes only. It does not run macros, evaluate formulas, fetch
linked resources, write originals, or use a language model to infer missing text.
"""

from __future__ import annotations

import csv
import posixpath
import re
import zipfile
from dataclasses import dataclass
from html.parser import HTMLParser
from io import BytesIO, StringIO
from pathlib import Path, PurePosixPath
from xml.etree import ElementTree as ET

from .file_policy import (
    HTML_EXTENSIONS,
    SUPPORTED_EXTENSIONS,
    TEXT_EXTENSIONS,
    credential_content,
    sensitive_filename,
)

MAX_FILE_BYTES = 15 * 1024 * 1024
MAX_CHARACTERS = 200_000
MAX_PAGES = 50
MAX_ROWS = 1_000
MAX_CELLS = 20_000
MAX_ZIP_MEMBERS = 1_000
MAX_ZIP_BYTES = 64 * 1024 * 1024
MAX_ZIP_MEMBER_BYTES = 16 * 1024 * 1024
IMAGE_EXTENSIONS = frozenset({".png", ".jpg", ".jpeg", ".webp", ".tif", ".tiff", ".bmp", ".heic"})
NS = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
}


@dataclass(frozen=True)
class Block:
    locator: str
    text: str


@dataclass
class Extracted:
    status: str
    blocks: list[Block]
    warnings: list[str]
    extractor: str


class _LimitReached(Exception):
    pass


class _InvalidArchive(Exception):
    pass


class _Collector:
    def __init__(self, extractor: str):
        self.extractor = extractor
        self.blocks: list[Block] = []
        self.warnings: list[str] = []
        self.characters = 0

    def warn(self, reason: str) -> None:
        if reason not in self.warnings:
            self.warnings.append(reason)

    def add(self, locator: str, text: str) -> bool:
        if not text or not text.strip():
            return True
        remaining = MAX_CHARACTERS - self.characters
        if remaining <= 0:
            self.warn("character_limit_reached")
            return False
        if len(text) > remaining:
            self.blocks.append(Block(locator, text[:remaining]))
            self.characters += remaining
            self.warn("character_limit_reached")
            return False
        self.blocks.append(Block(locator, text))
        self.characters += len(text)
        return True

    def finish(self, status: str | None = None) -> Extracted:
        return Extracted(
            status or ("partial" if self.warnings else "ok"),
            self.blocks,
            self.warnings,
            self.extractor,
        )


def _checked_zip(content: bytes) -> zipfile.ZipFile:
    archive = zipfile.ZipFile(BytesIO(content))
    infos = archive.infolist()
    if len(infos) > MAX_ZIP_MEMBERS:
        archive.close()
        raise _LimitReached
    total = 0
    names = set()
    for member in infos:
        name = member.filename
        path = PurePosixPath(name)
        if (
            name in names
            or path.is_absolute()
            or ".." in path.parts
            or "\\" in name
            or member.flag_bits & 1
        ):
            archive.close()
            raise _InvalidArchive
        names.add(name)
        total += member.file_size
        if member.file_size > MAX_ZIP_MEMBER_BYTES or total > MAX_ZIP_BYTES:
            archive.close()
            raise _LimitReached
    return archive


def _xml(archive: zipfile.ZipFile, name: str) -> ET.Element:
    return ET.fromstring(archive.read(name))


def _decode(content: bytes, collector: _Collector) -> str:
    if content.startswith((b"\xff\xfe", b"\xfe\xff")):
        try:
            return content.decode("utf-16")
        except UnicodeDecodeError:
            collector.warn("invalid_text_encoding")
            return content.decode("utf-16", errors="replace")
    try:
        return content.decode("utf-8-sig")
    except UnicodeDecodeError:
        collector.warn("invalid_text_encoding")
        return content.decode("utf-8-sig", errors="replace")


def _column(index: int) -> str:
    result = ""
    while index:
        index, remainder = divmod(index - 1, 26)
        result = chr(65 + remainder) + result
    return result


def _text(content: bytes, extension: str) -> Extracted:
    collector = _Collector("delimited_text" if extension in {".csv", ".tsv"} else "plain_text")
    decoded = _decode(content, collector)
    if extension in {".csv", ".tsv"}:
        cells = 0
        reader = csv.reader(
            StringIO(decoded, newline=""), delimiter="\t" if extension == ".tsv" else ","
        )
        try:
            for row_number, row in enumerate(reader, 1):
                if row_number > MAX_ROWS:
                    collector.warn("row_limit_reached")
                    break
                for column, value in enumerate(row, 1):
                    cells += 1
                    if cells > MAX_CELLS:
                        collector.warn("cell_limit_reached")
                        return collector.finish()
                    if not collector.add(f"row {row_number}, {_column(column)}", value):
                        return collector.finish()
        except csv.Error:
            collector.warn("delimited_text_parse_failed")
        return collector.finish()
    for number, line in enumerate(StringIO(decoded), 1):
        if number > MAX_ROWS:
            collector.warn("row_limit_reached")
            break
        if not collector.add(f"line {number}", line.rstrip("\r\n")):
            break
    return collector.finish()


class _HTMLText(HTMLParser):
    def __init__(self, collector: _Collector):
        super().__init__(convert_charrefs=True)
        self.collector = collector
        self.hidden: list[str] = []
        self.number = 0

    def handle_starttag(self, tag, attrs):
        if tag in {"script", "style", "template", "noscript", "head"}:
            self.hidden.append(tag)

    def handle_endtag(self, tag):
        if tag in self.hidden:
            self.hidden = self.hidden[:self.hidden.index(tag)]

    def handle_data(self, data):
        if self.hidden or not data.strip():
            return
        self.number += 1
        if not self.collector.add(f"html text {self.number}", data.strip()):
            raise _LimitReached


def _html(content: bytes) -> Extracted:
    collector = _Collector("html_text")
    parser = _HTMLText(collector)
    try:
        parser.feed(_decode(content, collector))
        parser.close()
    except _LimitReached:
        pass
    return collector.finish()


def _pdf(content: bytes) -> Extracted:
    from pypdf import PdfReader

    collector = _Collector("pypdf")
    if not content.startswith(b"%PDF-"):
        return Extracted("failed", [], ["document_parse_failed"], "pypdf")
    reader = PdfReader(BytesIO(content), strict=False)
    if reader.is_encrypted and not reader.decrypt(""):
        return Extracted("failed", [], ["encrypted_pdf"], "pypdf")
    page_count = len(reader.pages)
    if page_count > MAX_PAGES:
        collector.warn("page_limit_reached")
    missing = []
    failures = 0
    for index in range(min(page_count, MAX_PAGES)):
        try:
            text = reader.pages[index].extract_text() or ""
        except Exception:
            failures += 1
            collector.warn("page_extraction_failed")
            continue
        if not text.strip():
            missing.append(str(index + 1))
            continue
        if not collector.add(f"page {index + 1}", text):
            break
    if missing:
        collector.warn("pages_without_text:" + ",".join(missing))
    if not collector.blocks:
        if failures:
            return collector.finish("failed")
        collector.warn("ocr_required")
        return collector.finish("needs_ocr")
    return collector.finish()


def _docx(content: bytes, archive: zipfile.ZipFile) -> Extracted:
    from docx import Document
    from docx.oxml.table import CT_Tbl
    from docx.oxml.text.paragraph import CT_P
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    collector = _Collector("python-docx")
    document = Document(BytesIO(content))
    paragraph_number = table_number = rows = cells = 0
    for child in document.element.body.iterchildren():
        if isinstance(child, CT_P):
            paragraph_number += 1
            if not collector.add(f"paragraph {paragraph_number}", Paragraph(child, document).text):
                return collector.finish()
        elif isinstance(child, CT_Tbl):
            table_number += 1
            table = Table(child, document)
            for row_number, row in enumerate(table.rows, 1):
                rows += 1
                if rows > MAX_ROWS:
                    collector.warn("row_limit_reached")
                    return collector.finish()
                for column_number, cell in enumerate(row.cells, 1):
                    cells += 1
                    if cells > MAX_CELLS:
                        collector.warn("cell_limit_reached")
                        return collector.finish()
                    if cell.tables:
                        collector.warn("nested_tables_not_extracted")
                    if not collector.add(
                        f"table {table_number}, row {row_number}, {_column(column_number)}",
                        cell.text,
                    ):
                        return collector.finish()
    # Headers/footers are part of the document and can hold important context.
    for name in sorted(archive.namelist()):
        if re.fullmatch(r"word/(?:header|footer)\d+\.xml", name):
            xml = _xml(archive, name)
            for index, paragraph in enumerate(xml.findall(".//w:p", NS), 1):
                text = "".join(node.text or "" for node in paragraph.findall(".//w:t", NS))
                if not collector.add(f"{Path(name).stem}, paragraph {index}", text):
                    return collector.finish()
        if name in {"word/footnotes.xml", "word/endnotes.xml", "word/comments.xml"}:
            xml = _xml(archive, name)
            if any((node.text or "").strip() for node in xml.findall(".//w:t", NS)):
                collector.warn("notes_or_comments_not_extracted")
    if document.element.body.findall(".//w:txbxContent", NS):
        collector.warn("text_boxes_not_extracted")
    if document.element.body.findall(".//w:drawing", NS) or document.element.body.findall(
        ".//w:pict", NS
    ):
        collector.warn("images_not_ocr")
    return collector.finish()


def _xlsx(content: bytes, extension: str) -> Extracted:
    from openpyxl import load_workbook

    collector = _Collector("openpyxl")
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
    cached = None
    try:
        cached = load_workbook(BytesIO(content), read_only=True, data_only=True, keep_links=False)
        if extension == ".xlsm":
            collector.warn("macros_not_executed")
        rows_seen = cells_seen = 0
        for sheet in workbook.worksheets:
            # Missing dimensions must never trigger an unbounded preliminary scan.
            if sheet.max_row is None or sheet.max_column is None:
                collector.warn("worksheet_dimensions_missing")
            declared_rows = sheet.max_row if sheet.max_row is not None else MAX_ROWS
            declared_columns = sheet.max_column if sheet.max_column is not None else MAX_CELLS
            max_rows = min(declared_rows, MAX_ROWS - rows_seen)
            if declared_rows > max_rows:
                collector.warn("row_limit_reached")
            if max_rows <= 0:
                continue
            remaining = MAX_CELLS - cells_seen
            max_columns = min(declared_columns, remaining)
            if declared_columns > max_columns or max_rows * max_columns > remaining:
                collector.warn("cell_limit_reached")
            if max_columns <= 0:
                break
            cached_rows = cached[sheet.title].iter_rows(
                min_row=1, max_row=max_rows, max_col=max_columns
            )
            for row, cached_row in zip(
                sheet.iter_rows(min_row=1, max_row=max_rows, max_col=max_columns),
                cached_rows,
                strict=True,
            ):
                rows_seen += 1
                for cell, cached_cell in zip(row, cached_row, strict=True):
                    cells_seen += 1
                    if cells_seen > MAX_CELLS:
                        collector.warn("cell_limit_reached")
                        return collector.finish()
                    if cell.value is None:
                        continue
                    text = str(cell.value)
                    if cell.data_type == "f":
                        collector.warn("formulas_not_recalculated")
                        if cached_cell.value is None:
                            collector.warn("formula_cache_missing")
                            text += " [formula; cached value unavailable]"
                        else:
                            text += f" [formula; cached value: {cached_cell.value}]"
                    coordinate = f"'{sheet.title.replace(chr(39), chr(39) * 2)}'!{cell.coordinate}"
                    if not collector.add(coordinate, text):
                        return collector.finish()
        return collector.finish()
    finally:
        workbook.close()
        if cached is not None:
            cached.close()


def _pptx(archive: zipfile.ZipFile) -> Extracted:
    collector = _Collector("pptx_xml")
    presentation = _xml(archive, "ppt/presentation.xml")
    relationships = _xml(archive, "ppt/_rels/presentation.xml.rels")
    targets = {}
    for relation in relationships:
        if relation.get("Type", "").endswith("/slide") and relation.get("TargetMode") != "External":
            target = relation.get("Target", "")
            member = (
                target.lstrip("/")
                if target.startswith("/")
                else posixpath.normpath("ppt/" + target)
            )
            if not member.startswith("ppt/slides/") or ".." in PurePosixPath(member).parts:
                raise _InvalidArchive
            targets[relation.get("Id")] = member
    slide_ids = presentation.findall("p:sldIdLst/p:sldId", NS)
    if len(slide_ids) > MAX_PAGES:
        collector.warn("page_limit_reached")
    for number, slide_id in enumerate(slide_ids[:MAX_PAGES], 1):
        target = targets[slide_id.get(f"{{{NS['r']}}}id")]
        slide = _xml(archive, target)
        for index, paragraph in enumerate(slide.findall(".//a:p", NS), 1):
            text = "".join(node.text or "" for node in paragraph.findall(".//a:t", NS))
            if not collector.add(f"slide {number}, text {index}", text):
                return collector.finish()
        if slide.findall(".//p:pic", NS):
            collector.warn("images_not_ocr")
        if slide.findall(".//a:graphicData", NS):
            collector.warn("graphics_may_contain_unextracted_data")
    return collector.finish()


def _extract(filename: str, content: bytes) -> Extracted:
    """Extract bounded, located text; warnings never include document contents."""
    extension = Path(filename).suffix.lower()
    if extension in IMAGE_EXTENSIONS:
        return Extracted("needs_ocr", [], ["ocr_required"], "none")
    if extension not in SUPPORTED_EXTENSIONS:
        return Extracted("unsupported", [], ["unsupported_format"], "none")
    if len(content) > MAX_FILE_BYTES:
        return Extracted("partial", [], ["file_size_limit_reached"], "none")
    extractor = {
        ".pdf": "pypdf",
        ".docx": "python-docx",
        ".xlsx": "openpyxl",
        ".xlsm": "openpyxl",
        ".pptx": "pptx_xml",
    }.get(extension, "delimited_text" if extension in {".csv", ".tsv"} else "plain_text")
    try:
        if extension in TEXT_EXTENSIONS:
            return _text(content, extension)
        if extension in HTML_EXTENSIONS:
            return _html(content)
        if extension == ".pdf":
            return _pdf(content)
        with _checked_zip(content) as archive:
            if extension == ".docx":
                return _docx(content, archive)
            if extension in {".xlsx", ".xlsm"}:
                return _xlsx(content, extension)
            return _pptx(archive)
    except _LimitReached:
        return Extracted("partial", [], ["archive_limit_reached"], extractor)
    except Exception:
        return Extracted("failed", [], ["document_parse_failed"], extractor)


def extract(filename: str, content: bytes) -> Extracted:
    """Exclude detectable credentials from previews, model inputs and reports."""
    if sensitive_filename(filename):
        return Extracted("blocked", [], ["sensitive_file_blocked"], "none")
    # Inspect text before parser limits can hide a credential further down the file.
    extension = Path(filename).suffix.lower()
    if extension in TEXT_EXTENSIONS | HTML_EXTENSIONS:
        decoded = _decode(content, _Collector("none"))
        if credential_content(decoded):
            return Extracted("blocked", [], ["credential_content_detected"], "none")
    result = _extract(filename, content)
    if credential_content("\n".join(block.text for block in result.blocks)):
        return Extracted("blocked", [], ["credential_content_detected"], "none")
    return result
