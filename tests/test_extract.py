import zipfile
from io import BytesIO

from docx import Document
from openpyxl import Workbook
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from glocal_agent import extract as extraction
from glocal_agent.extract import extract


def pdf_bytes(texts):
    writer = PdfWriter()
    for text in texts:
        page = writer.add_blank_page(width=200, height=200)
        if text:
            font = DictionaryObject(
                {
                    NameObject("/Type"): NameObject("/Font"),
                    NameObject("/Subtype"): NameObject("/Type1"),
                    NameObject("/BaseFont"): NameObject("/Helvetica"),
                }
            )
            page[NameObject("/Resources")] = DictionaryObject(
                {
                    NameObject("/Font"): DictionaryObject({NameObject("/F1"): font}),
                }
            )
            stream = DecodedStreamObject()
            stream.set_data(f"BT /F1 12 Tf 20 100 Td ({text}) Tj ET".encode())
            page[NameObject("/Contents")] = stream
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


def archive_bytes(members):
    output = BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, content in members.items():
            archive.writestr(name, content)
    return output.getvalue()


def test_pdf_text_is_located_and_missing_pages_are_explicit():
    result = extract("invoice.pdf", pdf_bytes(["Invoice total 42", None]))
    assert result.status == "partial"
    assert result.blocks[0].locator == "page 1"
    assert "Invoice total 42" in result.blocks[0].text
    assert "pages_without_text:2" in result.warnings
    empty = extract("scan.pdf", pdf_bytes([None]))
    assert empty.status == "needs_ocr"
    assert empty.blocks == []


def test_pdf_page_limit(monkeypatch):
    monkeypatch.setattr(extraction, "MAX_PAGES", 1)
    result = extract("report.pdf", pdf_bytes(["one", "two"]))
    assert result.status == "partial"
    assert "page_limit_reached" in result.warnings
    assert [block.locator for block in result.blocks] == ["page 1"]


def test_encrypted_pdf():
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    writer.encrypt("private-password")
    output = BytesIO()
    writer.write(output)
    result = extract("private.pdf", output.getvalue())
    assert result.status == "failed"
    assert result.warnings == ["encrypted_pdf"]


def test_docx_preserves_order_and_locates_tables_and_headers():
    document = Document()
    document.add_paragraph("Purchase contract")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Supplier"
    table.cell(0, 1).text = "Glocal"
    document.add_paragraph("Signed today")
    document.sections[0].header.paragraphs[0].text = "Company header"
    output = BytesIO()
    document.save(output)
    result = extract("contract.docx", output.getvalue())
    assert result.status == "ok"
    assert [block.text for block in result.blocks[:4]] == [
        "Purchase contract",
        "Supplier",
        "Glocal",
        "Signed today",
    ]
    assert result.blocks[1].locator == "table 1, row 1, A"
    assert any(
        block.locator.startswith("header") and block.text == "Company header"
        for block in result.blocks
    )


def test_xlsx_preserves_formulas_without_inventing_values():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Prices"
    sheet["A1"] = "Item"
    sheet["B2"] = 21
    sheet["B3"] = 21
    sheet["B4"] = "=SUM(B2:B3)"
    output = BytesIO()
    workbook.save(output)
    result = extract("prices.xlsx", output.getvalue())
    assert result.status == "partial"
    assert {block.locator for block in result.blocks} == {
        "'Prices'!A1",
        "'Prices'!B2",
        "'Prices'!B3",
        "'Prices'!B4",
    }
    formula = next(block for block in result.blocks if block.locator.endswith("B4"))
    assert formula.text == "=SUM(B2:B3) [formula; cached value unavailable]"
    assert "formulas_not_recalculated" in result.warnings
    assert "formula_cache_missing" in result.warnings


def test_xlsx_cached_formula_value_is_labeled_as_cached():
    workbook = Workbook()
    workbook.active["A1"] = "=1+1"
    output = BytesIO()
    workbook.save(output)
    with zipfile.ZipFile(BytesIO(output.getvalue())) as archive:
        members = {name: archive.read(name) for name in archive.namelist()}
    members["xl/worksheets/sheet1.xml"] = members["xl/worksheets/sheet1.xml"].replace(
        b"<v></v>",
        b"<v>2</v>",
    )
    result = extract("formula.xlsx", archive_bytes(members))
    assert result.blocks[0].text == "=1+1 [formula; cached value: 2]"
    assert "formulas_not_recalculated" in result.warnings
    assert "formula_cache_missing" not in result.warnings


def test_xlsx_cell_and_row_limits(monkeypatch):
    workbook = Workbook()
    sheet = workbook.active
    for row in range(1, 4):
        sheet.cell(row, 1, "first")
        sheet.cell(row, 2, "second")
    output = BytesIO()
    workbook.save(output)
    monkeypatch.setattr(extraction, "MAX_ROWS", 2)
    monkeypatch.setattr(extraction, "MAX_CELLS", 3)
    result = extract("table.xlsx", output.getvalue())
    assert result.status == "partial"
    assert "row_limit_reached" in result.warnings
    assert "cell_limit_reached" in result.warnings
    assert len(result.blocks) == 3


def test_missing_xlsx_dimensions_are_bounded_without_full_scan(monkeypatch):
    workbook = Workbook()
    workbook.active["A1"] = "first"
    workbook.active["A100"] = "outside limit"
    output = BytesIO()
    workbook.save(output)
    with zipfile.ZipFile(BytesIO(output.getvalue())) as archive:
        members = {name: archive.read(name) for name in archive.namelist()}
    members["xl/worksheets/sheet1.xml"] = members["xl/worksheets/sheet1.xml"].replace(
        b'<dimension ref="A1:A100"/>',
        b"",
    )
    monkeypatch.setattr(extraction, "MAX_ROWS", 2)
    monkeypatch.setattr(extraction, "MAX_CELLS", 4)
    result = extract("no-dimensions.xlsx", archive_bytes(members))
    assert result.status == "partial"
    assert "worksheet_dimensions_missing" in result.warnings
    assert "outside limit" not in " ".join(block.text for block in result.blocks)


def test_pptx_follows_native_slide_order():
    members = {
        "ppt/presentation.xml": """<p:presentation
            xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"
            xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
            <p:sldIdLst><p:sldId id="256" r:id="rId2"/><p:sldId id="257" r:id="rId1"/>
            </p:sldIdLst></p:presentation>""",
        "ppt/_rels/presentation.xml.rels": """<Relationships
            xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
            <Relationship Id="rId1"
            Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide"
            Target="slides/slide1.xml"/>
            <Relationship Id="rId2"
            Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide"
            Target="slides/slide2.xml"/>
            </Relationships>""",
    }
    for number in (1, 2):
        members[f"ppt/slides/slide{number}.xml"] = f"""<p:sld
            xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"
            xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
            <p:cSld><p:spTree><p:sp><p:txBody><a:p><a:r><a:t>Original slide {number}</a:t>
            </a:r></a:p></p:txBody></p:sp></p:spTree></p:cSld></p:sld>"""
    result = extract("slides.pptx", archive_bytes(members))
    assert result.status == "ok"
    assert [(block.locator, block.text) for block in result.blocks] == [
        ("slide 1, text 1", "Original slide 2"),
        ("slide 2, text 1", "Original slide 1"),
    ]


def test_delimited_text_quotes_addresses_and_unicode():
    result = extract("table.csv", '\ufeffname,note\n张三,"hello, world"\n'.encode())
    assert result.status == "ok"
    assert [(block.locator, block.text) for block in result.blocks] == [
        ("row 1, A", "name"),
        ("row 1, B", "note"),
        ("row 2, A", "张三"),
        ("row 2, B", "hello, world"),
    ]
    utf16 = extract("notes.txt", "办公室".encode("utf-16"))
    assert utf16.status == "ok"
    assert utf16.blocks[0].text == "办公室"


def test_character_and_line_limits_are_explicit(monkeypatch):
    monkeypatch.setattr(extraction, "MAX_CHARACTERS", 5)
    result = extract("notes.md", b"123456789")
    assert result.status == "partial"
    assert result.blocks[0].text == "12345"
    assert "character_limit_reached" in result.warnings
    monkeypatch.setattr(extraction, "MAX_CHARACTERS", 100)
    monkeypatch.setattr(extraction, "MAX_ROWS", 2)
    result = extract("notes.txt", b"one\ntwo\nthree")
    assert result.status == "partial"
    assert "row_limit_reached" in result.warnings
    assert len(result.blocks) == 2


def test_archive_limits_and_invalid_members(monkeypatch):
    monkeypatch.setattr(extraction, "MAX_ZIP_MEMBER_BYTES", 16)
    result = extract("large.docx", archive_bytes({"word/document.xml": "x" * 1000}))
    assert result.status == "partial"
    assert result.warnings == ["archive_limit_reached"]
    monkeypatch.setattr(extraction, "MAX_ZIP_MEMBER_BYTES", 1024)
    result = extract("unsafe.docx", archive_bytes({"../private": "value"}))
    assert result.status == "failed"
    assert result.blocks == []


def test_unsupported_images_and_corruption_have_fixed_reasons():
    assert extract("photo.png", b"image").status == "needs_ocr"
    assert extract("legacy.xls", b"data").status == "unsupported"
    assert extract("legacy.doc", b"data").status == "unsupported"
    result = extract("broken.pdf", b"PRIVATE_SENTENCE_DO_NOT_REPORT")
    assert result.status == "failed"
    assert result.warnings == ["document_parse_failed"]
    assert result.blocks == []
    bad_zip = extract("broken.xlsx", b"PRIVATE_SENTENCE_DO_NOT_REPORT")
    assert bad_zip.warnings == ["document_parse_failed"]
